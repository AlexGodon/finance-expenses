import os

from openpyxl import load_workbook, Workbook

from reporting.month_tab import MonthTabWriter
from reporting.year_summary import YearSummaryWriter


class ExcelReporter:
    """Top-level facade. Opens/creates the workbook, runs the per-month writer for
    each month in the input df, then runs the year summary writer, then saves."""

    def __init__(self, file_path):
        self.file_path = file_path

    def write(self, df, category_map=None, budget_entries=None, fresh_file=False):
        os.makedirs(os.path.dirname(self.file_path) or '.', exist_ok=True)

        df = df.copy()
        df['month'] = df['date'].dt.to_period('M')

        wb = self._open_workbook(fresh_file=fresh_file)

        month_writer = MonthTabWriter(wb, category_map=category_map)
        year_writer = YearSummaryWriter(wb, category_map=category_map)

        months = sorted(df['month'].unique())
        data_ranges = {}
        for month in months:
            month_df = df[df['month'] == month].copy()
            data_ranges[month] = month_writer.write(month, month_df, budget_entries)

        year_writer.write(df, data_ranges)

        wb.save(self.file_path)

    def _open_workbook(self, fresh_file=False):
        if not fresh_file and os.path.exists(self.file_path):
            return load_workbook(self.file_path)
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']
        return wb
