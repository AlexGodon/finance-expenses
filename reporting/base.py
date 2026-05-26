import pandas as pd
from openpyxl.styles import Font

from helpers import col_range, month_tab_name, DATE_DISPLAY_FORMAT


BOLD = Font(bold=True)
BOLD_LARGE = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True, size=11)

TXN_COLUMNS = ['date', 'description', 'amount', 'category', 'subcategory', 'source']

MONTH_TAB_COLUMN_WIDTHS = {
    'A': 14, 'B': 55, 'C': 12, 'D': 18, 'E': 25, 'F': 12,
}
YEAR_TAB_COLUMN_WIDTHS = {
    'A': 30, 'B': 15, 'G': 60, 'H': 35,
}


def find_duplicates(df):
    """Find suspected cross-account duplicate pairs within a month.
    Matches transactions from different sources with equal absolute amounts
    and opposite signs (one positive, one negative)."""
    positives = df[df['amount'] > 0].copy().reset_index(drop=True)
    negatives = df[df['amount'] < 0].copy().reset_index(drop=True)

    used_pos = set()
    used_neg = set()
    pairs = []

    for i, pos_row in positives.iterrows():
        if i in used_pos:
            continue
        for j, neg_row in negatives.iterrows():
            if j in used_neg:
                continue
            if pos_row['source'] == neg_row['source']:
                continue
            if abs(pos_row['amount']) == abs(neg_row['amount']):
                pairs.append(pos_row)
                pairs.append(neg_row)
                used_pos.add(i)
                used_neg.add(j)
                break

    if not pairs:
        return pd.DataFrame(columns=df.columns)

    result = pd.DataFrame(pairs)
    result = result.sort_values(by='amount', key=lambda s: s.abs(), ascending=False)
    return result.reset_index(drop=True)


def build_month_summary(month_df, category_map):
    """Group transactions by category/subcategory and inject zero rows for known
    subcategories that have no transactions, so the report is structurally consistent."""
    summary = (
        month_df.groupby(['category', 'subcategory'])['amount']
        .sum()
        .reset_index()
    )
    categories_with_data = set(summary['category'].unique())

    if category_map:
        full_rows = []
        for cat, subcats in category_map.items():
            if cat not in categories_with_data:
                continue
            for sub in subcats:
                full_rows.append({'category': cat, 'subcategory': sub, 'amount': 0.0})
        if full_rows:
            full_df = pd.DataFrame(full_rows)
            summary = pd.concat([summary, full_df], ignore_index=True)
            summary = (
                summary.groupby(['category', 'subcategory'])['amount']
                .sum()
                .reset_index()
            )

    if category_map:
        order_map = {cat: i for i, cat in enumerate(category_map.keys())}
        summary = summary.sort_values(
            'category', key=lambda col: col.map(
                lambda c: order_map.get(c, len(order_map))
            )
        )
    else:
        summary = summary.sort_values('category')
    return summary


def read_all_month_tabs(wb):
    """Read transaction data from all existing monthly tabs in the workbook."""
    all_dfs = []

    for tab_name in wb.sheetnames:
        if tab_name.isdigit():
            continue
        ws = wb[tab_name]
        if ws.max_row < 3:
            continue
        headers = [ws.cell(row=2, column=c).value for c in range(1, 7)]
        if headers != TXN_COLUMNS:
            continue

        rows = []
        for r in range(3, ws.max_row + 1):
            first_cell = ws.cell(row=r, column=1).value
            if first_cell is None or first_cell == 'Suspected Duplicates':
                break
            rows.append([ws.cell(row=r, column=c).value for c in range(1, 7)])

        if rows:
            month_df = pd.DataFrame(rows, columns=TXN_COLUMNS)
            month_df['date'] = pd.to_datetime(month_df['date'], format=DATE_DISPLAY_FORMAT)
            month_df['amount'] = pd.to_numeric(month_df['amount'], errors='coerce').fillna(0.0)
            all_dfs.append(month_df)

    if not all_dfs:
        return pd.DataFrame(columns=TXN_COLUMNS)
    return pd.concat(all_dfs, ignore_index=True)


def get_month_data_range(wb, month):
    """Return (first_row, last_row) of transaction data on a month tab."""
    tab = month_tab_name(month)
    if tab not in wb.sheetnames:
        return 3, 3
    ws = wb[tab]
    last_row = 2
    for r in range(3, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if val is None or val == 'Suspected Duplicates':
            break
        last_row = r
    return 3, last_row


def build_budget_sumifs(entry, data_start, data_end):
    """Build the SUMIFS formula string for a single budget YAML entry's column C."""
    formula = entry.get('formula', {})
    ftype = formula.get('type', 'none')

    if ftype == 'none':
        return 0

    rc = col_range('C', data_start, data_end)
    rd = col_range('D', data_start, data_end)
    re_ = col_range('E', data_start, data_end)
    rf = col_range('F', data_start, data_end)

    if ftype in ('category', 'subcategory'):
        cat = formula.get('category', '')
        sub = formula.get('subcategory')
        if sub:
            return f'=SUMIFS({rc},{rd},"{cat}",{re_},"{sub}")'
        return f'=SUMIFS({rc},{rd},"{cat}")'

    if ftype == 'lookups':
        parts = []
        for lk in formula.get('lookups', []):
            cat = lk.get('category', '')
            sub = lk.get('subcategory')
            if sub:
                parts.append(f'SUMIFS({rc},{rd},"{cat}",{re_},"{sub}")')
            else:
                parts.append(f'SUMIFS({rc},{rd},"{cat}")')
        return '=' + '+'.join(parts) if parts else 0

    if ftype == 'source_negatives':
        sources = formula.get('sources', [])
        parts = [f'SUMIFS({rc},{rf},"{src}",{rc},"<"&0)' for src in sources]
        return '=' + '+'.join(parts)

    return 0


class BaseTabWriter:
    """Shared state and utilities for tab writers (month tab + year summary).
    Subclasses implement `write(...)` with their specific arguments."""

    def __init__(self, wb, category_map=None):
        self.wb = wb
        self.category_map = category_map or {}

    def _replace_sheet(self, sheet_name, index=None):
        """Delete existing sheet of this name (if any) and create a fresh one."""
        if sheet_name in self.wb.sheetnames:
            del self.wb[sheet_name]
        if index is None:
            return self.wb.create_sheet(sheet_name)
        return self.wb.create_sheet(sheet_name, index)

    def _apply_column_widths(self, ws, widths):
        for letter, width in widths.items():
            ws.column_dimensions[letter].width = width
