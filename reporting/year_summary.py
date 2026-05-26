import pandas as pd
from openpyxl.worksheet.properties import Outline

from helpers import col_range, month_tab_name, month_full_name
from reporting.base import (
    BaseTabWriter,
    BOLD,
    BOLD_LARGE,
    YEAR_TAB_COLUMN_WIDTHS,
    build_month_summary,
    read_all_month_tabs,
    get_month_data_range,
)


CSV_INSTRUCTIONS = [
    ("How to get monthly CSV files:", BOLD_LARGE),
    ("Wealthsimple:", BOLD),
    ("Click Settings (bottom left) > Accounts > Select Chequing >", None),
    ("Documents > Monthly statements > Download CSV of desired month", None),
    ("ScotiaBank:", BOLD),
    ("Click Credit Card > In transactions > Set Date Range > Download as CSV", None),
    ("CIBC:", BOLD),
    ("Left side menu bar > Download Transactions > Select Account >", None),
    ("Set Date Range > Download Transactions", None),
]

VENV_INSTRUCTIONS = [
    ("Activate virtual environment:", BOLD),
    ("source .venv/bin/activate", None),
]


class YearSummaryWriter(BaseTabWriter):
    """Writes the per-year summary tabs (one sheet per calendar year, named '2026' etc.).
    Pulls data for current months from the in-memory df and preserves prior months by
    reading their existing tabs."""

    def write(self, df, data_ranges):
        current_months = set(df['month'].unique())

        existing = read_all_month_tabs(self.wb)
        if not existing.empty:
            existing['month'] = existing['date'].dt.to_period('M')
            preserved = existing[~existing['month'].isin(current_months)]
            combined = pd.concat([df, preserved], ignore_index=True)
        else:
            combined = df.copy()

        years = sorted(combined['month'].apply(lambda m: m.year).unique())

        for year in years:
            year_months = sorted(
                [m for m in combined['month'].unique() if m.year == year],
                reverse=True,
            )
            self._write_year_sheet(year, year_months, combined, data_ranges)

    def _write_year_sheet(self, year, year_months, combined, data_ranges):
        ws = self._replace_sheet(str(year), index=0)
        ws.sheet_properties.outlinePr = Outline(summaryBelow=False)

        row_cursor = 1
        ws.cell(row=row_cursor, column=1, value=f'Grand Total {year}').font = BOLD_LARGE
        row_cursor += 1

        row_cursor = self._write_interest_line(ws, year_months, data_ranges, row_cursor)
        row_cursor += 1

        for month in year_months:
            dr = data_ranges.get(month) or get_month_data_range(self.wb, month)
            month_df = combined[combined['month'] == month]
            summary = build_month_summary(month_df, self.category_map)
            row_cursor = self._write_month_block(ws, month, summary, row_cursor, dr)

        self._apply_column_widths(ws, YEAR_TAB_COLUMN_WIDTHS)
        self._write_instructions(ws)

    def _write_interest_line(self, ws, year_months, data_ranges, row_cursor):
        ws.cell(row=row_cursor, column=1, value='Interest Earned (YTD):').font = BOLD
        interest_parts = []
        for m in year_months:
            tab = month_tab_name(m)
            dr = data_ranges.get(m) or get_month_data_range(self.wb, m)
            rc = col_range('C', dr[0], dr[1])
            rd = col_range('D', dr[0], dr[1])
            interest_parts.append(
                f'SUMIFS(\'{tab}\'!{rc},\'{tab}\'!{rd},"interest-earned")'
            )
        ws.cell(row=row_cursor, column=2,
                value='=' + '+'.join(interest_parts)).font = BOLD
        return row_cursor + 1

    def _write_month_block(self, ws, month, summary, start_row, data_range):
        row_cursor = start_row
        tab = month_tab_name(month)
        ref = f"'{tab}'"
        ds, de = data_range
        rc = col_range('C', ds, de)
        rd = col_range('D', ds, de)
        re_ = col_range('E', ds, de)

        ws.cell(row=row_cursor, column=1, value=month_full_name(month)).font = BOLD_LARGE
        row_cursor += 1

        ws.cell(row=row_cursor, column=1, value='Total Income:')
        ws.cell(row=row_cursor, column=2,
                value=f'=SUMIFS({ref}!{rc},{ref}!{rc},">"&0)')
        income_row = row_cursor
        row_cursor += 1

        ws.cell(row=row_cursor, column=1, value='Total Expenses:')
        ws.cell(row=row_cursor, column=2,
                value=f'=SUMIFS({ref}!{rc},{ref}!{rc},"<"&0)')
        row_cursor += 1

        ws.cell(row=row_cursor, column=1, value='Net:')
        ws.cell(row=row_cursor, column=2,
                value=f'=B{income_row}+B{income_row + 1}').font = BOLD
        row_cursor += 2

        current_cat = None
        cat_header_row = None
        cat_first_row = None
        for _, srow in summary.iterrows():
            if srow['category'] != current_cat:
                if current_cat is not None:
                    ws.cell(row=cat_header_row, column=2,
                            value=f"=SUM(B{cat_first_row}:B{row_cursor - 1})").font = BOLD
                    ws.row_dimensions.group(cat_first_row, row_cursor - 1,
                                            outline_level=1, hidden=True)

                current_cat = srow['category']
                ws.cell(row=row_cursor, column=1, value=current_cat.capitalize()).font = BOLD
                cat_header_row = row_cursor
                row_cursor += 1
                cat_first_row = row_cursor

            ws.cell(row=row_cursor, column=1, value=f"  {srow['subcategory']}")
            ws.cell(row=row_cursor, column=2,
                    value=f'=SUMIFS({ref}!{rc},{ref}!{rd},"{srow["category"]}",{ref}!{re_},"{srow["subcategory"]}")')
            row_cursor += 1

        if current_cat is not None:
            ws.cell(row=cat_header_row, column=2,
                    value=f"=SUM(B{cat_first_row}:B{row_cursor - 1})").font = BOLD
            ws.row_dimensions.group(cat_first_row, row_cursor - 1,
                                    outline_level=1, hidden=True)

        return row_cursor + 2

    def _write_instructions(self, ws):
        for i, (text, font) in enumerate(CSV_INSTRUCTIONS, start=1):
            cell = ws.cell(row=i, column=7, value=text)
            if font:
                cell.font = font

        for i, (text, font) in enumerate(VENV_INSTRUCTIONS, start=1):
            cell = ws.cell(row=i, column=8, value=text)
            if font:
                cell.font = font
