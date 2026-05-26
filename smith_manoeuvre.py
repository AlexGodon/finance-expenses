#!/usr/bin/env python3
"""Smith Manoeuvre Mortgage Calculator.

Simulates the Smith Manoeuvre month-by-month until the mortgage principal
reaches $0. Outputs a console summary and an Excel file (.xlsx) with live
formulas.

All rates/percentages are entered as human-readable numbers (e.g. 4.30 for
4.30%). The program handles conversion to decimals internally.
"""

import argparse
import math
import os

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# ── CLI ─────────────────────────────────────────────────────────────────────

def parse_args():
    p = argparse.ArgumentParser(
        description="Smith Manoeuvre Mortgage Calculator — Canadian semi-annual compounding",
    )
    p.add_argument("--initial-amount", type=float, required=True,
                   help="Mortgage principal in dollars")
    p.add_argument("--amortization", type=int, required=True,
                   help="Amortization period in years")
    p.add_argument("--mortgage-rate", type=float, required=True,
                   help="Annual mortgage rate (e.g. 4.30 for 4.30%%)")
    p.add_argument("--heloc-rate", type=float, required=True,
                   help="Annual HELOC rate (e.g. 5.45 for 5.45%%)")
    p.add_argument("--tax-bracket", type=float, required=True,
                   help="Marginal tax bracket (e.g. 27 for 27%%)")
    p.add_argument("--rental-income", type=float, required=True,
                   help="Monthly rental income in dollars")
    p.add_argument("--frequency", default="monthly",
                   choices=["monthly", "biweekly", "weekly"],
                   help="Payment frequency (default: monthly)")
    p.add_argument("--output", default="smith_manoeuvre_output.xlsx",
                   help="Output Excel filename (default: smith_manoeuvre_output.xlsx)")
    return p.parse_args()


# ── Calculations ────────────────────────────────────────────────────────────

def calc_monthly_mortgage_rate(annual_rate_pct):
    """Canadian semi-annual compounding -> equivalent monthly rate."""
    r = annual_rate_pct / 100
    return (1 + r / 2) ** (1 / 6) - 1


def calc_mortgage_payment(principal, annual_rate_pct, amortization_years,
                          frequency="monthly"):
    """Compute the periodic mortgage payment."""
    annual_rate = annual_rate_pct / 100

    if frequency == "monthly":
        rate = (1 + annual_rate / 2) ** (1 / 6) - 1
        n = amortization_years * 12
    elif frequency == "biweekly":
        rate = (1 + annual_rate / 2) ** (1 / 13) - 1
        n = amortization_years * 26
    elif frequency == "weekly":
        rate = (1 + annual_rate / 2) ** (1 / 26) - 1
        n = amortization_years * 52
    else:
        raise ValueError(f"Unknown frequency: {frequency}")

    return principal * rate / (1 - (1 + rate) ** (-n))


def monthly_equivalent_payment(principal, annual_rate_pct, amortization_years,
                               frequency):
    """Return the monthly-equivalent payment for any frequency."""
    if frequency == "monthly":
        return calc_mortgage_payment(principal, annual_rate_pct,
                                     amortization_years, "monthly")
    per_period = calc_mortgage_payment(principal, annual_rate_pct,
                                       amortization_years, frequency)
    periods_per_year = 26 if frequency == "biweekly" else 52
    return per_period * periods_per_year / 12


def calc_yearly_heloc_payment(heloc_balance, monthly_heloc_rate,
                              pull_from_heloc):
    """Closed-form D — rounded up to nearest dollar.

    D = ceil(B * r + P * (1 - (1+r)^(-12)))
    Sized so month-12 HELOC interest is fully covered.
    """
    d = (heloc_balance * monthly_heloc_rate
         + pull_from_heloc * (1 - (1 + monthly_heloc_rate) ** (-12)))
    return math.ceil(d)


def calc_total_original_interest(principal, annual_rate_pct,
                                 amortization_years, frequency):
    """Total interest over the full original amortization (no SM)."""
    payment = monthly_equivalent_payment(principal, annual_rate_pct,
                                         amortization_years, frequency)
    rate = calc_monthly_mortgage_rate(annual_rate_pct)
    balance = principal
    total = 0.0
    for _ in range(amortization_years * 12):
        interest = balance * rate
        total += interest
        balance -= (payment - interest)
        if balance <= 0:
            break
    return total


# ── Simulation ──────────────────────────────────────────────────────────────

def simulate(initial_amount, amortization_years, mortgage_rate_pct,
             heloc_rate_pct, tax_bracket_pct, rental_income,
             frequency="monthly"):
    """Month-by-month Smith Manoeuvre simulation until mortgage = 0."""
    monthly_mtg_rate = calc_monthly_mortgage_rate(mortgage_rate_pct)
    monthly_heloc_rate = (heloc_rate_pct / 100) / 12
    tax_bracket = tax_bracket_pct / 100
    total_payment = monthly_equivalent_payment(
        initial_amount, mortgage_rate_pct, amortization_years, frequency,
    )

    rows = []
    new_balance = initial_amount
    original_balance = initial_amount
    heloc_loan = 0.0
    heloc_room = 0.0
    prev_full_pct = 0.0
    D = 0

    month_num = 0
    max_months = amortization_years * 12 * 2

    while new_balance > 0 and month_num < max_months:
        month_num += 1
        month_in_year = ((month_num - 1) % 12) + 1
        year = (month_num - 1) // 12 + 1

        if month_in_year == 1:
            D = calc_yearly_heloc_payment(heloc_loan, monthly_heloc_rate,
                                          rental_income)

        # ── Mortgage ──
        inter_payment = new_balance * monthly_mtg_rate
        princ_payment = total_payment - inter_payment

        original_inter = original_balance * monthly_mtg_rate
        original_princ = total_payment - original_inter
        original_balance = max(original_balance - original_princ, 0)

        ending_balance = new_balance - princ_payment

        extra_princ = rental_income - D
        extra_princ = min(extra_princ, max(ending_balance, 0))

        new_balance = max(ending_balance - extra_princ, 0)

        # ── HELOC ──
        heloc_room += princ_payment + extra_princ

        if month_num == 1:
            heloc_loan = rental_income
        else:
            heloc_loan = heloc_loan + prev_full_pct + rental_income - D

        full_pct = heloc_loan * monthly_heloc_rate
        tax_refund = full_pct * tax_bracket
        pct_inter_paid = full_pct - tax_refund
        heloc_left = heloc_room - heloc_loan

        rows.append({
            "year": year,
            "month": month_in_year,
            "global_month": month_num,
            "princ_payment": princ_payment,
            "inter_payment": inter_payment,
            "original_inter": original_inter,
            "total_payment": total_payment,
            "ending_balance": ending_balance,
            "extra_princ": extra_princ,
            "new_balance": new_balance,
            "heloc_room": heloc_room,
            "heloc_left": heloc_left,
            "heloc_loan": heloc_loan,
            "full_pct_heloc": full_pct,
            "tax_refund": tax_refund,
            "pct_inter_paid": pct_inter_paid,
            "pay_heloc_directly": D,
        })

        prev_full_pct = full_pct
        if new_balance <= 0:
            break

    return rows, total_payment


# ── Console output ──────────────────────────────────────────────────────────

def print_summary(rows, args, total_payment):
    monthly_mtg_rate = calc_monthly_mortgage_rate(args.mortgage_rate)
    monthly_heloc_rate = (args.heloc_rate / 100) / 12
    real_heloc = args.heloc_rate * (1 - args.tax_bracket / 100)

    print("=" * 64)
    print("  Smith Manoeuvre Mortgage Calculator")
    print("=" * 64)

    print("\n  Input Values:")
    print(f"    Initial Amount:      ${args.initial_amount:,.2f}")
    print(f"    Amortization:        {args.amortization} years")
    print(f"    Mortgage Rate:       {args.mortgage_rate:.2f}%")
    print(f"    HELOC Rate:          {args.heloc_rate:.2f}%")
    print(f"    Tax Bracket:         {args.tax_bracket:.2f}%")
    print(f"    Rental Income:       ${args.rental_income:,.2f}/month")
    print(f"    Frequency:           {args.frequency.capitalize()}")

    print("\n  Calculated Values:")
    print(f"    Monthly Mtg Rate:    {monthly_mtg_rate * 100:.4f}%")
    print(f"    Monthly HELOC Rate:  {monthly_heloc_rate * 100:.4f}%")
    print(f"    Monthly Payment:     ${total_payment:,.2f}")
    print(f"    Real HELOC Rate:     {real_heloc:.2f}%")

    years = {}
    for r in rows:
        y = r["year"]
        if y not in years:
            years[y] = dict(inter_sum=0, extra_sum=0, heloc_inter_sum=0,
                            tax_refund_sum=0, end_balance=0, heloc_loan=0,
                            D=r["pay_heloc_directly"])
        years[y]["inter_sum"] += r["inter_payment"]
        years[y]["extra_sum"] += r["extra_princ"]
        years[y]["heloc_inter_sum"] += r["full_pct_heloc"]
        years[y]["tax_refund_sum"] += r["tax_refund"]
        years[y]["end_balance"] = r["new_balance"]
        years[y]["heloc_loan"] = r["heloc_loan"]

    hdr = (f"  {'Yr':>3}  {'New Balance':>14}  {'HELOC Loan':>14}"
           f"  {'Mtg Interest':>14}  {'Extra Princ':>14}  {'Pay HELOC D':>12}")
    sep = (f"  {'---':>3}  {'-' * 14}  {'-' * 14}"
           f"  {'-' * 14}  {'-' * 14}  {'-' * 12}")

    print(f"\n  Year-by-Year Summary:\n{hdr}\n{sep}")
    for y in sorted(years):
        d = years[y]
        print(f"  {y:>3}  ${d['end_balance']:>13,.2f}  ${d['heloc_loan']:>13,.2f}"
              f"  ${d['inter_sum']:>13,.2f}  ${d['extra_sum']:>13,.2f}"
              f"  ${d['D']:>11,}")

    total_months = len(rows)
    yy, mm = divmod(total_months, 12)
    total_mtg_interest = sum(r["inter_payment"] for r in rows)
    total_orig_interest = calc_total_original_interest(
        args.initial_amount, args.mortgage_rate, args.amortization,
        args.frequency,
    )
    interest_saved = total_orig_interest - total_mtg_interest
    final_heloc = rows[-1]["heloc_loan"]
    total_heloc_interest = sum(r["full_pct_heloc"] for r in rows)
    total_tax_refund = sum(r["tax_refund"] for r in rows)

    print(f"\n  Final Results:")
    print(f"    Mortgage paid off in:         {yy} years, {mm} months")
    print(f"    Total orig. interest (no SM): ${total_orig_interest:>12,.2f}")
    print(f"    Total mtg interest (w/ SM):   ${total_mtg_interest:>12,.2f}")
    print(f"    Mortgage interest saved:      ${interest_saved:>12,.2f}")
    print(f"    Final HELOC balance:          ${final_heloc:>12,.2f}")
    print(f"    Total HELOC interest:         ${total_heloc_interest:>12,.2f}")
    print(f"    Total tax refunds:            ${total_tax_refund:>12,.2f}")
    print(f"    Net HELOC cost:               "
          f"${total_heloc_interest - total_tax_refund:>12,.2f}")
    net = interest_saved - (total_heloc_interest - total_tax_refund)
    print(f"    Net overall savings:          ${net:>12,.2f}")


# ── Excel output ────────────────────────────────────────────────────────────

def write_excel(rows, args, total_payment, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Smith Manoeuvre"

    bold = Font(bold=True)
    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4",
                           fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF")
    curr = '#,##0.00'
    pct = '0.00%'
    pct6 = '0.000000%'

    # ── Static values (rows 1-4) ──
    ws.merge_cells('A1:G1')
    ws['A1'] = 'Smith Manoeuvre Mortgage Calculator'
    ws['A1'].font = Font(bold=True, size=14)

    for i, lbl in enumerate(['', 'Initial Amount', 'Amortization (yr)',
                             'Mortgage Rate', 'HELOC Rate', 'Tax Bracket',
                             'Rental Income']):
        ws.cell(row=3, column=i + 1, value=lbl).font = bold

    ws['B4'] = args.initial_amount;  ws['B4'].number_format = curr
    ws['C4'] = args.amortization
    ws['D4'] = args.mortgage_rate / 100;  ws['D4'].number_format = pct
    ws['E4'] = args.heloc_rate / 100;     ws['E4'].number_format = pct
    ws['F4'] = args.tax_bracket / 100;    ws['F4'].number_format = pct
    ws['G4'] = args.rental_income;        ws['G4'].number_format = curr

    # ── Derived values (rows 6-7) ──
    for i, lbl in enumerate(['', 'Mo. Mtg Rate', 'Mo. HELOC Rate',
                             'Mo. Payment', 'Real HELOC Rate']):
        ws.cell(row=6, column=i + 1, value=lbl).font = bold

    ws['B7'] = '=((1+D4/2)^(1/6))-1';       ws['B7'].number_format = pct6
    ws['C7'] = '=E4/12';                      ws['C7'].number_format = pct6
    ws['D7'] = '=B4*B7/(1-(1+B7)^(-C4*12))'; ws['D7'].number_format = curr
    ws['E7'] = '=E4*(1-F4)';                  ws['E7'].number_format = pct

    # ── Column headers (row 9) ──
    DATA_HEADERS = [
        'Period', 'Princ. Payment', 'Inter Payment', 'Original Inter Paym.',
        'Total Payment', 'Ending Balance', 'Extra Princ. Paym.',
        'New Balance', 'HELOC Room', 'HELOC Left', 'HELOC Loan',
        'Full % HELOC', 'Tax Refund', '% Inter. Paid',
        'Pay HELOC D', 'Orig. Balance',
    ]
    for i, h in enumerate(DATA_HEADERS):
        c = ws.cell(row=9, column=i + 1, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center')

    # ── Data rows ──
    cur = 10
    prev = None
    current_year = 0
    year_first = {}
    year_last = {}
    year_summary = {}

    for idx, rd in enumerate(rows):
        yr, mo, gm = rd['year'], rd['month'], rd['global_month']
        is_last = (idx == len(rows) - 1)

        if yr != current_year:
            ws.cell(row=cur, column=1, value=f"Year {yr}").font = bold
            cur += 1
            current_year = yr
            year_first[yr] = cur

        r = cur
        ws.cell(row=r, column=1, value=f"Month {mo}")

        # O (15): Pay HELOC Directly
        if mo == 1:
            if yr == 1:
                f15 = '=CEILING($G$4*(1-(1+$C$7)^(-12)),1)'
            else:
                pl = year_last[yr - 1]
                f15 = f'=CEILING(K{pl}*$C$7+$G$4*(1-(1+$C$7)^(-12)),1)'
        else:
            f15 = f'=O{year_first[yr]}'
        ws.cell(row=r, column=15).value = f15

        # C (3): Inter Payment
        if gm == 1:
            ws.cell(row=r, column=3).value = '=$B$4*$B$7'
        else:
            ws.cell(row=r, column=3).value = f'=H{prev}*$B$7'

        # B (2): Princ. Payment
        ws.cell(row=r, column=2).value = f'=$D$7-C{r}'

        # D (4): Original Inter Paym.
        if gm == 1:
            ws.cell(row=r, column=4).value = '=$B$4*$B$7'
        else:
            ws.cell(row=r, column=4).value = f'=P{prev}*$B$7'

        # E (5): Total Payment
        ws.cell(row=r, column=5).value = '=$D$7'

        # F (6): Ending Balance
        if gm == 1:
            ws.cell(row=r, column=6).value = f'=$B$4-B{r}'
        else:
            ws.cell(row=r, column=6).value = f'=H{prev}-B{r}'

        # G (7): Extra Princ. Paym.
        ws.cell(row=r, column=7).value = f'=MIN($G$4-O{r},MAX(F{r},0))'

        # H (8): New Balance
        ws.cell(row=r, column=8).value = f'=MAX(F{r}-G{r},0)'

        # I (9): HELOC Room
        if gm == 1:
            ws.cell(row=r, column=9).value = f'=B{r}+G{r}'
        else:
            ws.cell(row=r, column=9).value = f'=I{prev}+B{r}+G{r}'

        # K (11): HELOC Loan
        if gm == 1:
            ws.cell(row=r, column=11).value = '=$G$4'
        else:
            ws.cell(row=r, column=11).value = f'=K{prev}+L{prev}+$G$4-O{r}'

        # L (12): Full % HELOC
        ws.cell(row=r, column=12).value = f'=K{r}*$C$7'

        # M (13): Tax Refund
        ws.cell(row=r, column=13).value = f'=L{r}*$F$4'

        # N (14): % Inter. Paid
        ws.cell(row=r, column=14).value = f'=L{r}-M{r}'

        # J (10): HELOC Left
        ws.cell(row=r, column=10).value = f'=I{r}-K{r}'

        # P (16): Original Balance (helper)
        if gm == 1:
            ws.cell(row=r, column=16).value = f'=$B$4*(1+$B$7)-$D$7'
        else:
            ws.cell(row=r, column=16).value = f'=P{prev}*(1+$B$7)-$D$7'

        for col in range(2, 17):
            ws.cell(row=r, column=col).number_format = curr

        year_last[yr] = r
        prev = r
        cur += 1

        # Year-end summary
        if mo == 12 or is_last:
            sr = cur
            first = year_first[yr]
            last = r
            ws.cell(row=sr, column=1,
                    value=f"W/ SM Year {yr} Totals").font = bold
            for col in [2, 3, 4, 5, 7, 12, 13, 14]:
                cl = get_column_letter(col)
                ws.cell(row=sr, column=col).value = (
                    f'=SUM({cl}{first}:{cl}{last})')
                ws.cell(row=sr, column=col).number_format = curr
                ws.cell(row=sr, column=col).font = bold
            for col in [6, 8, 9, 10, 11, 15]:
                cl = get_column_letter(col)
                ws.cell(row=sr, column=col).value = f'={cl}{last}'
                ws.cell(row=sr, column=col).number_format = curr
                ws.cell(row=sr, column=col).font = bold
            year_summary[yr] = sr
            cur += 2

    # ── Full summary (mirrors console output) ──
    monthly_mtg_rate = calc_monthly_mortgage_rate(args.mortgage_rate)
    monthly_heloc_rate = (args.heloc_rate / 100) / 12
    real_heloc = args.heloc_rate * (1 - args.tax_bracket / 100)
    total_months = len(rows)
    yy, mm = divmod(total_months, 12)
    total_mtg_interest = sum(r["inter_payment"] for r in rows)
    total_orig_interest = calc_total_original_interest(
        args.initial_amount, args.mortgage_rate, args.amortization,
        args.frequency,
    )
    interest_saved = total_orig_interest - total_mtg_interest
    final_heloc = rows[-1]["heloc_loan"]
    total_heloc_interest = sum(r["full_pct_heloc"] for r in rows)
    total_tax_refund = sum(r["tax_refund"] for r in rows)
    net_heloc_cost = total_heloc_interest - total_tax_refund
    net_savings = interest_saved - net_heloc_cost

    section_font = Font(bold=True, size=12)
    label_font = Font(bold=True)
    sum_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3",
                           fill_type="solid")

    fr = cur + 1

    # ── Input Values ──
    ws.cell(row=fr, column=1, value="Input Values").font = section_font
    fr += 1
    input_rows = [
        ("Initial Amount", args.initial_amount, curr),
        ("Amortization", f"{args.amortization} years", None),
        ("Mortgage Rate", args.mortgage_rate / 100, pct),
        ("HELOC Rate", args.heloc_rate / 100, pct),
        ("Tax Bracket", args.tax_bracket / 100, pct),
        ("Rental Income", args.rental_income, curr),
        ("Frequency", args.frequency.capitalize(), None),
    ]
    for lbl, val, fmt in input_rows:
        ws.cell(row=fr, column=1, value=lbl).font = label_font
        c = ws.cell(row=fr, column=2, value=val)
        if fmt:
            c.number_format = fmt
        fr += 1

    fr += 1  # blank row

    # ── Calculated Values ──
    ws.cell(row=fr, column=1, value="Calculated Values").font = section_font
    fr += 1
    calc_rows = [
        ("Monthly Mortgage Rate", monthly_mtg_rate, pct6),
        ("Monthly HELOC Rate", monthly_heloc_rate, pct6),
        ("Monthly Payment", total_payment, curr),
        ("Real HELOC Rate (after tax)", real_heloc / 100, pct),
    ]
    for lbl, val, fmt in calc_rows:
        ws.cell(row=fr, column=1, value=lbl).font = label_font
        c = ws.cell(row=fr, column=2, value=val)
        c.number_format = fmt
        fr += 1

    fr += 1  # blank row

    # ── Year-by-Year Summary table ──
    ws.cell(row=fr, column=1, value="Year-by-Year Summary").font = section_font
    fr += 1

    yby_headers = ['Year', 'New Balance', 'HELOC Loan', 'Mtg Interest',
                   'Extra Princ', 'Pay HELOC D']
    for i, h in enumerate(yby_headers):
        c = ws.cell(row=fr, column=i + 1, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center')
    fr += 1

    year_data = {}
    for r in rows:
        y = r["year"]
        if y not in year_data:
            year_data[y] = dict(inter_sum=0, extra_sum=0, end_balance=0,
                                heloc_loan=0, D=r["pay_heloc_directly"])
        year_data[y]["inter_sum"] += r["inter_payment"]
        year_data[y]["extra_sum"] += r["extra_princ"]
        year_data[y]["end_balance"] = r["new_balance"]
        year_data[y]["heloc_loan"] = r["heloc_loan"]

    for y in sorted(year_data):
        d = year_data[y]
        ws.cell(row=fr, column=1, value=y)
        ws.cell(row=fr, column=2, value=round(d['end_balance'], 2))
        ws.cell(row=fr, column=3, value=round(d['heloc_loan'], 2))
        ws.cell(row=fr, column=4, value=round(d['inter_sum'], 2))
        ws.cell(row=fr, column=5, value=round(d['extra_sum'], 2))
        ws.cell(row=fr, column=6, value=d['D'])
        for col in range(2, 6):
            ws.cell(row=fr, column=col).number_format = curr
        ws.cell(row=fr, column=6).number_format = '#,##0'
        fr += 1

    fr += 1  # blank row

    # ── Final Results ──
    ws.cell(row=fr, column=1, value="Final Results").font = section_font
    fr += 1
    final_rows = [
        ("Mortgage paid off in", f"{yy} years, {mm} months", None),
        ("Total orig. interest (no SM)", total_orig_interest, curr),
        ("Total mtg interest (w/ SM)", total_mtg_interest, curr),
        ("Mortgage interest saved", interest_saved, curr),
        ("Final HELOC balance", final_heloc, curr),
        ("Total HELOC interest", total_heloc_interest, curr),
        ("Total tax refunds", total_tax_refund, curr),
        ("Net HELOC cost", net_heloc_cost, curr),
        ("Net overall savings", net_savings, curr),
    ]
    for lbl, val, fmt in final_rows:
        ws.cell(row=fr, column=1, value=lbl).font = label_font
        c = ws.cell(row=fr, column=2)
        if fmt:
            c.value = round(val, 2)
            c.number_format = fmt
        else:
            c.value = val
        fr += 1

    # Highlight the net savings row
    ws.cell(row=fr - 1, column=1).fill = sum_fill
    ws.cell(row=fr - 1, column=2).fill = sum_fill

    # Column widths
    ws.column_dimensions['A'].width = 24
    for col in range(2, 17):
        ws.column_dimensions[get_column_letter(col)].width = 18

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)
    return output_path


# ── Main ────────────────────────────────────────────────────────────────────

def main():
    args = parse_args()

    rows, total_payment = simulate(
        initial_amount=args.initial_amount,
        amortization_years=args.amortization,
        mortgage_rate_pct=args.mortgage_rate,
        heloc_rate_pct=args.heloc_rate,
        tax_bracket_pct=args.tax_bracket,
        rental_income=args.rental_income,
        frequency=args.frequency,
    )

    print_summary(rows, args, total_payment)

    out = write_excel(rows, args, total_payment, args.output)
    print(f"\n  Excel file written to: {out}")


if __name__ == "__main__":
    main()
